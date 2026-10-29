"""Load and version the team's Alert_Logic workbook. Parameters are found by the label text in
Config column B wherever they sit (teams re-letter and re-order sections), never by cell position.
Loading never stops a run: every problem becomes an issue with its cell, and a safe value is used –
from the last good logic version if given, else Ver 1's value."""
import difflib
import re
import shutil
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

PARAM_LABELS: dict[str, tuple[str, str]] = {
    "gmv_pace_low": ("A", "GMV pace thấp nếu <"),
    "spend_pace_low": ("A", "Spend pace thấp nếu <"),
    "spend_pace_high": ("A", "Spend pace quá cao nếu >"),
    "tacos_over": ("A", "TACoS vượt target nếu"),
    "not_running": ("A", '"Không chạy" nếu Spend pace <'),
    "budget_cap": ("A", "Ngưỡng cạn budget"),
    "trend_spend": ("B", "Spend"),
    "trend_acos": ("B", "ACoS"),
    "trend_tacos": ("B", "TACoS"),
    "trend_min_spend": ("B", "Sàn L7D Spend của PL"),
    "camp_cut": ("C", "CUT/OPTIMIZE nếu ACoS camp"),
    "camp_good": ("C", "MAINTAIN/PUSH hoặc SCALE UP nếu"),
    "camp_optimize": ("C", "OPTIMIZE nếu ACoS camp"),
    "min_clicks": ("C", "Sàn click tối thiểu"),
    "excess_high": ("C", "Excess spend để lên HIGH"),
    "action_window": ("D", 'Cửa sổ "đã action gần đây"'),
    "followup_days": ("D", "Hạn follow-up"),
    "improving_ratio": ("D", 'Coi là "đang cải thiện"'),
}
# Score model (team v3): weights, sensitivities, score cut-offs, guards.
V3_LABELS: dict[str, tuple[str, str]] = {
    "w1": ("B", "w1 –"), "w2": ("B", "w2 –"), "w3": ("B", "w3 –"), "w4": ("B", "w4 –"), "w5": ("B", "w5 –"),
    "w6": ("B", "w6 –"),
    "sens_tacos": ("C", "TACoS lệch so với target"), "sens_ads": ("C", "Ads dependency lệch"),
    "sens_trend": ("C", "Trend TACoS (±)"), "sens_gmv": ("C", "GMV pace lệch so với 1.0"),
    "sens_spend": ("C", "Spend pace lệch so với 1.0"),
    "score_sd": ("D", "SCALE DOWN nếu điểm"), "score_opt": ("D", "OPTIMIZE nếu điểm"), "score_up": ("D", "SCALE UP nếu điểm"),
    "sd_room": ("D", "Nới ngưỡng SCALE DOWN"), "sd_over": ("D", "Siết ngưỡng SCALE DOWN"),
    "min_time_gone": ("E", "Time gone tối thiểu"), "min_spend": ("E", "Spend MTD tối thiểu"),
}
# Inventory guard (Config K, logic v4) – optional: files without it use the defaults, silently.
INV_LABELS: dict[str, tuple[str, str]] = {
    "inv_low_days": ("K", 'Ngưỡng "sắp hết hàng"'), "inv_high_days": ("K", 'Ngưỡng "dư hàng"'),
    "inv_incoming_days": ("K", "Cộng hàng đang về trong"), "inv_sales_days": ("K", "Lượng bán trung bình tính trên"),
    "inv_attention": ("K", 'Nhóm "Cần chú ý"'), "inv_critical": ("K", 'Nhóm "Nghiêm trọng"'),
    "inv_excess": ("K", 'Nhóm "Dư hàng"'), "inv_camp_block": ("K", "Campaign bị chặn khi"),
    "inv_stale_days": ("K", "Số liệu tồn kho bị coi là quá cũ"),
}
DEFAULT_INV = {"inv_low_days": 14.0, "inv_high_days": 90.0, "inv_incoming_days": 7.0, "inv_sales_days": 7.0,
               "inv_attention": 0.2, "inv_critical": 0.5, "inv_excess": 0.5, "inv_camp_block": 0.5, "inv_stale_days": 2.0}
TEXT_LABELS = {"team_label": "Nhãn team được tracking"}
# Readable names where the file's label prefix is too short to understand on its own.
DISPLAY = {"w1": "Trọng số w1 – TACoS vs target riêng", "w2": "Trọng số w2 – TACoS vs target Department",
           "w3": "Trọng số w3 – phụ thuộc ads", "w4": "Trọng số w4 – trend TACoS L7D", "w5": "Trọng số w5 – GMV pace",
           "w6": "Trọng số w6 – Spend pace", "sens_tacos": "Độ nhạy: TACoS lệch so với target (±)",
           "sens_ads": "Độ nhạy: Ads dependency lệch benchmark (±)", "sens_trend": "Độ nhạy: Trend TACoS (±)",
           "sens_gmv": "Độ nhạy: GMV pace lệch 1.0 (±)", "sens_spend": "Độ nhạy: Spend pace lệch 1.0 (±)",
           "score_sd": "SCALE DOWN nếu điểm ≥", "score_opt": "OPTIMIZE nếu điểm ≥", "score_up": "SCALE UP nếu điểm ≤",
           "sd_room": "Nới ngưỡng SCALE DOWN khi CÒN ROOM", "sd_over": "Siết ngưỡng SCALE DOWN khi VƯỢT",
           "min_time_gone": "Time gone tối thiểu để chấm điểm", "min_spend": "Spend MTD tối thiểu để chấm điểm ($)"}
ALL_LABELS = {**PARAM_LABELS, **V3_LABELS, **INV_LABELS}
V1_ONLY = ("gmv_pace_low", "spend_pace_low", "spend_pace_high", "tacos_over")
DEFAULT_V3 = {"w1": 3.0, "w2": 3.0, "w3": 2.0, "w4": 2.0, "w5": 2.0, "w6": 2.0, "sens_tacos": 0.5, "sens_ads": 0.5,
              "sens_trend": 0.5, "sens_gmv": 0.3, "sens_spend": 0.5, "score_sd": 50.0, "score_opt": 25.0,
              "score_up": -25.0, "sd_room": 15.0, "sd_over": -10.0, "min_time_gone": 0.15, "min_spend": 50.0}
# Other wordings seen in team versions of the file.
ALIASES = {"trend_min_spend": ("Sàn L7D Spend",), "improving_ratio": ('"Đang cải thiện"',)}
# Bare metric names ("Spend") only count inside a section whose header mentions "trend".
_EXACT = {"trend_spend", "trend_acos", "trend_tacos"}
# Ver 1 values: the last-resort fallback.
DEFAULT_PARAMS = {"gmv_pace_low": 0.85, "spend_pace_low": 0.9, "spend_pace_high": 1.15, "tacos_over": 1.1,
                  "not_running": 0.1, "budget_cap": 1.0, "trend_spend": 0.2, "trend_acos": 0.2, "trend_tacos": 0.2,
                  "trend_min_spend": 50.0, "camp_cut": 1.15, "camp_good": 1.0, "camp_optimize": 1.3, "min_clicks": 15.0,
                  "excess_high": 30.0, "action_window": 3.0, "followup_days": 3.0, "improving_ratio": 0.9}
DEFAULT_K = {"SPS_PRODUCTS": 3.0, "SPS_BRANDS": 3.0, "SPS_DISPLAY": 5.0}
VALID_LABELS = ("REVIEW ADS", "REVIEW SKU", "SCALE")
MODELS = {"rules_v1": "Mô hình rule Ver 1 (pacing ngưỡng)", "score_v3": "Mô hình chấm điểm TACoS có trọng số (v3)"}
_DEFAULT_LABELS = {
    **{k: "REVIEW ADS" for k in ["U|U|U", "U|U|F", "U|U|D", "U|F|U", "F|U|F", "F|U|D", "D|U|F", "D|U|D"]},
    **{k: "REVIEW SKU" for k in ["U|D|U", "F|U|U", "F|F|U", "F|D|U", "D|U|U", "D|F|U", "D|D|U"]},
    **{k: "SCALE" for k in ["U|F|D", "U|D|F", "U|D|D", "F|D|D", "D|D|D"]},
}
DEFAULT_MATRIX = {f"{s}|{a}|{t}": (_DEFAULT_LABELS.get(f"{s}|{a}|{t}", ""), "")
                  for s in "UFD" for a in "UFD" for t in "UFD"}
# Sanity ranges (value outside → reported, value still used).
RANGES = {"trend_spend": (0.01, 2), "trend_acos": (0.01, 2), "trend_tacos": (0.01, 2), "min_clicks": (1, 1000),
          "action_window": (1, 60), "followup_days": (1, 60), "improving_ratio": (0.1, 1.5),
          "camp_cut": (0.5, 5), "camp_good": (0.1, 5), "camp_optimize": (0.5, 5), "gmv_pace_low": (0.1, 2),
          "spend_pace_low": (0.1, 2), "spend_pace_high": (0.5, 5), "tacos_over": (0.5, 5),
          **{f"w{i}": (0, 100) for i in range(1, 7)},
          **{k: (0.01, 10) for k in ("sens_tacos", "sens_ads", "sens_trend", "sens_gmv", "sens_spend")},
          "min_time_gone": (0, 1), "inv_low_days": (1, 365), "inv_high_days": (1, 3650), "inv_incoming_days": (0, 90),
          "inv_sales_days": (1, 90), **{k: (0, 1) for k in ("inv_attention", "inv_critical", "inv_excess", "inv_camp_block")},
          "inv_stale_days": (0, 60)}
_INFO_ROWS = ("config", "as-of", "đầu tháng", "số ngày", "time gone", "l7d", "p7d")


class LogicError(Exception):
    pass


@dataclass
class Logic:
    path: Path
    version: int
    params: dict[str, float]
    k_by_type: dict[str, float]
    matrix: dict[str, tuple[str, str]]
    tiers_text: str
    cells: dict[str, str] = field(default_factory=dict)  # key → "Config!C22" / "matrix:KEY" → "Matrix_PL!F5"
    issues: list[dict] = field(default_factory=list)
    model: str = "rules_v1"  # "score_v3" when the file has weights w1–w6
    texts: dict[str, str] = field(default_factory=dict)  # text parameters, e.g. team_label


def version_from_name(name: str) -> int:
    """'Alert_Logic_Ver 12' → 12, '…_Team_v3' → 3, no version → 1."""
    m = re.search(r"(?:^|[^a-z])v(?:er)?\s*(\d+)", Path(name).stem, re.IGNORECASE)
    return int(m.group(1)) if m else 1


def _parse(value, formula) -> tuple[float | None, str | None]:
    """Config value → (number, None) or (None, issue code). '1,25' and '110%' are accepted."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), None
    if value is None or str(value).strip() == "":
        return None, "FORMULA" if isinstance(formula, str) and formula.startswith("=") else "BLANK"
    text = str(value).strip().replace(" ", "")
    pct = text.endswith("%")
    try:
        return float(text.rstrip("%").replace(",", ".")) / (100 if pct else 1), None
    except ValueError:
        return None, "NOT_NUMBER"


def _key_for(label: str, section_title: str) -> str | None:
    low = label.lower()
    for key, prefix in TEXT_LABELS.items():
        if low.startswith(prefix.lower()):
            return key
    for key, (_, prefix) in ALL_LABELS.items():
        if key in _EXACT:
            if label == prefix and "trend" in section_title.lower():
                return key
            continue
        if any(low.startswith(p.lower()) for p in (prefix, *ALIASES.get(key, ()))):
            return key
    return None


def _scan(ws):
    """Yield (row, kind, key_or_label, value): kind 'param', 'k' or 'unused'. Section header rows are skipped."""
    title = ""
    for row in ws.iter_rows(min_col=2, max_col=3):
        label = str(row[0].value).strip() if row[0].value is not None else ""
        if not label:
            continue
        if re.match(r"^[A-Z]\.\s", label):
            title = label
            continue
        if label.upper().startswith("SPS_") or "k =" in title:
            yield row[0].row, "k", label, row[1].value
            continue
        key = _key_for(label, title)
        if key in TEXT_LABELS:
            yield row[0].row, "text", key, row[1].value
        elif key:
            yield row[0].row, "param", key, row[1].value
        elif "tự tính" in title.lower():
            continue  # computed block (e.g. Department totals): the app computes these itself
        elif title and not label.lower().startswith(_INFO_ROWS):
            yield row[0].row, "unused", label, row[1].value


_MESSAGES = {
    "BLANK": "ô đang trống",
    "FORMULA": "là công thức chưa có giá trị đã tính (file không được lưu bằng Excel)",
    "NOT_NUMBER": "không phải số",
    "MISSING": "không tìm thấy trong sheet Config",
    "DUPLICATE": "xuất hiện nhiều lần – dùng dòng đầu tiên",
    "OUT_OF_RANGE": "nằm ngoài khoảng hợp lý",
    "ORDER": "không hợp lý so với ngưỡng liên quan",
    "UNUSED": "app chưa dùng tham số này",
    "NO_CONFIG": "file không có sheet Config",
    "NO_MATRIX": "file không có sheet Matrix_PL",
    "MATRIX_KEY": "thiếu dòng trong Matrix_PL",
    "MATRIX_LABEL": "nhãn không hợp lệ (chỉ được REVIEW ADS / REVIEW SKU / SCALE / trống)",
}
SEVERITY_HINT = {"UNUSED": "info", "DUPLICATE": "warn", "OUT_OF_RANGE": "warn", "ORDER": "warn"}


def _issue(code, key, label, cell, value, used, used_from) -> dict:
    return {"code": code, "key": key, "label": label, "cell": cell, "value": None if value is None else str(value),
            "used": used, "used_from": used_from, "hint": SEVERITY_HINT.get(code, "error"),
            "message": f"{label or key} ({cell or 'Config'}): {_MESSAGES[code]}"}


def load_logic(path: Path, fallback: "Logic | None" = None) -> Logic:
    path = Path(path)
    fb_params = dict(fallback.params) if fallback else dict(DEFAULT_PARAMS)
    fb_k = dict(fallback.k_by_type) if fallback else dict(DEFAULT_K)
    fb_matrix = dict(fallback.matrix) if fallback else dict(DEFAULT_MATRIX)
    fb_name = fallback.path.name if fallback else "mặc định (Ver 1)"
    params, k_by_type, cells, issues, matrix, tiers, texts = {}, {}, {}, [], {}, [], {}
    wb = openpyxl.load_workbook(path, data_only=True)
    formulas = None  # loaded only if a value is missing, to tell blank from an unsaved formula
    try:
        if "Config" not in wb.sheetnames:
            issues.append(_issue("NO_CONFIG", None, "Config", None, None, None, fb_name))
        else:
            for row, kind, key, value in _scan(wb["Config"]):
                cell = f"Config!C{row}"
                if kind == "unused":
                    issues.append(_issue("UNUSED", None, key, cell, value, None, None))
                    continue
                if kind == "text":
                    cells[key] = cell
                    texts[key] = str(value or "").strip()
                    if not texts[key]:
                        issues.append(_issue("BLANK", key, TEXT_LABELS[key], cell, value, None, None))
                    continue
                target, fb = (params, fb_params) if kind == "param" else (k_by_type, fb_k)
                label = ALL_LABELS[key][1] if kind == "param" else f"k cho {key}"
                if key in target:
                    issues.append(_issue("DUPLICATE", key, label, cell, value, target[key], None))
                    continue
                if value is None and formulas is None:
                    formulas = openpyxl.load_workbook(path)["Config"]
                raw = formulas.cell(row=row, column=3).value if formulas is not None else None
                number, problem = _parse(value, raw)
                cells[key] = cell
                if problem:
                    used = fb.get(key, DEFAULT_PARAMS.get(key, DEFAULT_V3.get(key, DEFAULT_K.get(key))))
                    issues.append(_issue(problem, key, label, cell, value if problem == "NOT_NUMBER" else raw, used, fb_name))
                    number = used
                target[key] = number
        model = "score_v3" if any(f"w{i}" in params for i in range(1, 7)) else "rules_v1"
        required = [k for k in PARAM_LABELS if not (model == "score_v3" and k in V1_ONLY)]
        if model == "score_v3":
            required += list(V3_LABELS)
            if "team_label" not in texts:
                issues.append(_issue("MISSING", "team_label", TEXT_LABELS["team_label"], None, None, None, None))
        for key in required:
            if key not in params:
                params[key] = fb_params.get(key, DEFAULT_PARAMS.get(key, DEFAULT_V3.get(key)))
                issues.append(_issue("MISSING", key, ALL_LABELS[key][1], None, None, params[key], fb_name))
        if model == "rules_v1":
            for key in V3_LABELS:
                params.pop(key, None)
        for key, value in DEFAULT_INV.items():
            params.setdefault(key, fb_params.get(key, value))
        for key in DEFAULT_K:
            if key not in k_by_type:
                k_by_type[key] = fb_k.get(key, DEFAULT_K[key])
                issues.append(_issue("MISSING", key, f"k cho {key}", None, None, k_by_type[key], fb_name))
        for key, (lo, hi) in RANGES.items():
            if key in cells and key in params and not lo <= params[key] <= hi:
                issues.append(_issue("OUT_OF_RANGE", key, ALL_LABELS[key][1], cells[key], params[key], params[key], None)
                              | {"range": [lo, hi]})
        if params["camp_good"] > params["camp_cut"]:
            issues.append(_issue("ORDER", "camp_good", PARAM_LABELS["camp_good"][1], cells.get("camp_good"),
                                 params["camp_good"], params["camp_good"], None)
                          | {"detail": f"phải ≤ ngưỡng CUT ({params['camp_cut']})"})
        if model == "rules_v1" and params["spend_pace_low"] > params["spend_pace_high"]:
            issues.append(_issue("ORDER", "spend_pace_low", PARAM_LABELS["spend_pace_low"][1], cells.get("spend_pace_low"),
                                 params["spend_pace_low"], params["spend_pace_low"], None)
                          | {"detail": f"phải < Spend pace quá cao ({params['spend_pace_high']})"})

        msheet = next((n for n in ("Matrix_PL", "Matrix_Product Line") if n in wb.sheetnames), None)
        if msheet is None:
            matrix = fb_matrix
            issues.append(_issue("NO_MATRIX", None, "Matrix_PL", None, None, None, fb_name))
        else:
            for n, row in enumerate(wb[msheet].iter_rows(min_row=5, values_only=True), 5):
                if not (row and isinstance(row[0], str) and re.fullmatch(r"[UFD]\|[UFD]\|[UFD]", row[0].strip())):
                    continue
                k = row[0].strip()
                label = str(row[5] or "").strip() if len(row) > 5 else ""
                hint = str(row[6] or "") if len(row) > 6 else ""
                cells[f"matrix:{k}"] = f"{msheet}!F{n}"
                if label and label.upper() not in VALID_LABELS:
                    used = fb_matrix.get(k, ("", ""))[0]
                    guess = difflib.get_close_matches(label.upper(), VALID_LABELS, n=1)
                    issues.append(_issue("MATRIX_LABEL", f"matrix:{k}", f"Matrix_PL {k}", f"{msheet}!F{n}", label, used, fb_name)
                                  | {"suggest": guess[0] if guess else None})
                    label = used
                matrix[k] = (label.upper(), hint)
            for k in DEFAULT_MATRIX:
                if k not in matrix:
                    matrix[k] = fb_matrix.get(k, DEFAULT_MATRIX[k])
                    issues.append(_issue("MATRIX_KEY", f"matrix:{k}", f"Matrix_PL {k}", None, None, matrix[k][0], fb_name))
        tsheet = next((n for n in ("Logic_Tiers", "Logic Tier") if n in wb.sheetnames), None)
        if tsheet:
            for row in wb[tsheet].iter_rows(values_only=True):
                parts = [str(v) for v in row if v not in (None, "")]
                if parts:
                    tiers.append(" | ".join(parts))
    finally:
        wb.close()
    return Logic(path, version_from_name(path.name), params, k_by_type, matrix, "\n".join(tiers), cells, issues,
                 model, texts)


def write_logic_version(logic: Logic, changes: dict[str, float | str], note: str, out_dir: Path) -> Path:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    existing = [version_from_name(p.name) for p in out_dir.glob("*.xlsx")]
    n = max([logic.version, *existing]) + 1  # never reuse a number, even after switching back to an old version
    out = out_dir / f"Alert_Logic_Ver {n}.xlsx"
    if out.exists():
        raise ValueError(f"{out.name} đã tồn tại")
    shutil.copyfile(logic.path, out)
    wb = openpyxl.load_workbook(out)
    params = {k: v for k, v in changes.items() if not k.startswith("matrix:")}
    rows = {key: r for r, kind, key, _ in _scan(wb["Config"]) if kind == "param"}
    for key, value in params.items():
        if key not in rows:
            raise ValueError(f"File logic không có dòng cho tham số {key} – không thể ghi")
        wb["Config"].cell(row=rows[key], column=3, value=value)
    for key, label in changes.items():
        if key.startswith("matrix:"):
            mkey = key.split(":", 1)[1]
            for row in wb["Matrix_PL"].iter_rows(min_row=5):
                if row[0].value == mkey:
                    row[5].value = label or None
    if "README_Flow" not in wb.sheetnames:
        wb.create_sheet("README_Flow")
    wb["README_Flow"].append(["", f"Ver {n} – {date.today():%d/%m/%Y}", note])
    wb.save(out)
    return out


def valid_change(key: str, value, logic: Logic) -> bool:
    if key in logic.params:
        return isinstance(value, (int, float)) and not isinstance(value, bool)
    return key.startswith("matrix:") and key[7:] in logic.matrix and isinstance(value, str)
