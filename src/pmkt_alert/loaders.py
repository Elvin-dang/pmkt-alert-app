"""Detect the input workbooks by content and load them into normalized frames.
v4: one Data workbook (raw_sso_us, Product Info, KPI, INELIGIBLE ASIN, ads_camp, Summary, Inventory AMZ) + Activity
Log + Logic. Older sets (separate tracker / ads files, CSV tables, Defind) still load."""
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

from .xlsx import is_csv, read_sheet, read_sheet_with_rows, sheet_names, to_num

ROLES = ("tracker", "ads", "activity", "logic")
OPTIONAL_ROLES = ("defind",)
# The tracker workbook's tables; each may instead come as its own CSV.
TRACKER_SHEETS = {"raw": "raw_sso_us", "products": "Product Info", "kpi": "KPI", "ineligible": "INELIGIBLE ASIN"}
TRACKER_REQUIRED = ("raw", "products", "kpi")
CSV_COLUMNS = {  # a CSV's table is recognised by these header columns
    "raw": {"report_date", "asin", "sp_spend"}, "products": {"asin", "product_line_old", "status"},
    "kpi": {"Month", "ASIN", "Spend_final"}, "ineligible": {"asin", "delivery_status"},
    "ads": {"campaign_id", "spend_usd", "account_name"}, "activity": {"Level", "username"}, "defind": {"Tên File"},
}
ROLE_HINT = {
    "tracker": "file Data: cần sheet raw_sso_us, Product Info, KPI, ads_camp",
    "ads": "sheet ads_camp (trong file Data)",
    "activity": "file Activity Log: cần cột Level, username",
    "logic": "file Logic: cần sheet Config, Matrix_Product Line",
    "defind": "file Defind: cần cột Tên File",
}
BLANK = {"#N/A", "0", "0.0", "None", "nan"}


class InputError(Exception):
    """A structural problem the user must fix before a run can work. `location` points at it:
    {"file", "sheet", "columns" | "cell", "row"}."""

    def __init__(self, message: str, location: dict | None = None):
        super().__init__(message)
        self.location = location


@dataclass
class Inputs:
    raw: pd.DataFrame
    products: pd.DataFrame
    kpi: pd.DataFrame
    ineligible: pd.DataFrame
    ads: pd.DataFrame
    activity: pd.DataFrame
    event_days: dict[date, str] = field(default_factory=dict)
    defind_text: str = ""
    sources: dict = field(default_factory=dict)  # name → {"file", "sheet", "letters"} for evidence
    inventory: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=["date", "asin", "sku", "stock"]))
    contrib: dict = field(default_factory=dict)  # date → {"gmv", "ppc"} share of the month (Summary sheet)


def _first_row(path: Path) -> list[str]:
    if is_csv(path):
        df = read_sheet(path)
        return list(df.columns)
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        row = next(ws.iter_rows(max_row=1, values_only=True), ())
    finally:
        wb.close()
    return [str(v).strip() for v in row if v is not None]


def role_of(path: Path) -> str | None:
    """Workbook → tracker / ads / logic / activity / defind; CSV → its table (raw, products, kpi, ineligible, ads…)."""
    if is_csv(path):
        head = set(_first_row(path))
        return next((t for t, cols in CSV_COLUMNS.items() if cols <= head), None)
    sheets = set(sheet_names(path))
    if {"raw_sso_us", "Product Info", "KPI"} <= sheets:
        return "tracker"
    if "ads_camp" in sheets:
        return "ads"
    if "Config" in sheets and any(n.startswith("Matrix") for n in sheets):
        return "logic"
    head = _first_row(path)
    if "Level" in head and "username" in head:
        return "activity"
    if "Tên File" in head:
        return "defind"
    return None


def input_files(paths: list[Path]) -> list[Path]:
    files: list[Path] = []
    for p in map(Path, paths):
        files += sorted([*p.glob("*.xlsx"), *p.glob("*.csv")]) if p.is_dir() else [p]
    return [f for f in files if not f.name.startswith("~$")]


def detect_inputs(paths: list[Path]) -> dict[str, Path]:
    """role/table → file. The tracker is one workbook, or its tables as CSVs (raw, products, kpi[, ineligible]);
    a CSV table also overrides the matching sheet of a tracker workbook."""
    files = input_files(paths)
    found: dict[str, list[Path]] = {}
    for f in files:
        try:
            role = role_of(f)
        except Exception:
            continue
        if role:
            found.setdefault(role, []).append(f)
    if len(found.get("activity", [])) > 1:  # daily exports pile up in Downloads: use the newest
        found["activity"] = [max(found["activity"], key=lambda f: f.stat().st_mtime)]
    if "ads" not in found and len(found.get("tracker", [])) == 1 and "ads_camp" in sheet_names(found["tracker"][0]):
        found["ads"] = list(found["tracker"])  # v4: the Data workbook carries ads_camp too
    if "tracker" not in found and any(t in found for t in TRACKER_SHEETS):
        missing = [TRACKER_SHEETS[t] for t in TRACKER_REQUIRED if t not in found]
        if missing:
            raise InputError("Thiếu file Tracking dạng CSV: " + ", ".join(missing)
                             + " (hoặc dùng 1 file .xlsx có các sheet raw_sso_us, Product Info, KPI)")
    needed = [r for r in ROLES if not (r == "tracker" and all(t in found for t in TRACKER_REQUIRED))]
    for role in needed + [t for t in TRACKER_SHEETS if t in found]:
        if role not in found:
            raise InputError(f"Thiếu {ROLE_HINT[role]}")
        if len(found[role]) > 1:
            name = ROLE_HINT[role].split(':')[0] if role in ROLE_HINT else f"file {TRACKER_SHEETS[role]}"
            raise InputError(f"Có nhiều hơn một {name}: " + ", ".join(f.name for f in found[role]))
    return {role: files_[0] for role, files_ in found.items()
            if role in needed or role in TRACKER_SHEETS or (role in OPTIONAL_ROLES and len(files_) == 1)}


def _text_dates(s: pd.Series) -> pd.Series:
    """Date text from CSVs, vectorised: ISO, or d/m/y vs m/d/y decided from the values (default day-first)."""
    iso = pd.to_datetime(s, errors="coerce", format="ISO8601")
    if iso.notna().mean() >= 0.9:
        return iso
    parts = s.str.extract(r"^\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})")
    a, b = pd.to_numeric(parts[0], errors="coerce"), pd.to_numeric(parts[1], errors="coerce")
    dayfirst = bool((a > 12).any()) or not bool((b > 12).any())
    day, month = (a, b) if dayfirst else (b, a)
    year = pd.to_numeric(parts[2], errors="coerce")
    year = year.where(year >= 100, year + 2000)
    out = pd.to_datetime(pd.DataFrame({"year": year, "month": month, "day": day}), errors="coerce")
    return out.where(out.notna(), iso)


def _to_date(s: pd.Series) -> pd.Series:
    texts = s.map(lambda v: isinstance(v, str))
    if texts.any():
        parsed = _text_dates(s[texts].astype(str))
        out = s.copy().astype(object)
        out[texts] = [None if pd.isna(v) else v.date() for v in parsed]
        rest = ~texts
        if rest.any():
            out[rest] = _to_date_values(s[rest])
        return out
    return _to_date_values(s)


def _to_date_values(s: pd.Series) -> pd.Series:
    def one(v):
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        if isinstance(v, (int, float)) and not pd.isna(v):
            return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))).date()
        t = pd.to_datetime(v, errors="coerce")
        return None if pd.isna(t) else t.date()
    return s.map(one)


def _clean_text(s: pd.Series) -> pd.Series:
    return s.map(lambda v: "" if v is None or str(v).strip() in BLANK else str(v).strip())


def _need(df: pd.DataFrame, cols: list[str], where: str, path: Path | None = None, sheet: str | None = None) -> None:
    missing = [c for c in cols if c not in df.columns]
    if missing:
        file = Path(path).name if path else None
        raise InputError(f"{where} thiếu cột: {', '.join(missing)} – thêm cột (đúng tên) vào dòng tiêu đề (dòng 1)"
                         + (f" của {file} › {sheet or where}" if file else ""),
                         {"file": file, "sheet": sheet or where, "columns": missing, "row": 1})


def _contributions(path: Path) -> dict:
    """Summary sheet: date → {"gmv": %Contribution GMV, "ppc": %Contribution PPC} (daily phasing of the month)."""
    if "Summary" not in sheet_names(path):
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Summary"]
        ws.reset_dimensions()
        rows = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
    finally:
        wb.close()
    label = lambda r: str(r[3]).strip().lower() if len(r) > 3 and r[3] is not None else ""  # noqa: E731
    days = next((r for r in rows if label(r) == "day"), None)
    gmv = next((r for r in rows if label(r) == "%contribution gmv"), None)
    ppc = next((r for r in rows if label(r) == "%contribution ppc"), None)
    if not days or not gmv:
        return {}
    out = {}
    for i in range(4, len(days)):
        d = days[i]
        if isinstance(d, (datetime, date)):
            num = lambda r: float(r[i]) if r and i < len(r) and isinstance(r[i], (int, float)) else 0.0  # noqa: E731
            out[d.date() if isinstance(d, datetime) else d] = {"gmv": num(gmv), "ppc": num(ppc) if ppc else num(gmv)}
    return out


def _campaign_id(name: str) -> str:
    """New ads_camp exports have no campaign_id: Amazon names end with it ("…_1782098480")."""
    tail = str(name).strip().rsplit("_", 1)[-1]
    return tail if tail.isdigit() else str(name).strip()


def _event_days(path: Path) -> dict[date, str]:
    if "Summary" not in sheet_names(path):
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Summary"]
        ws.reset_dimensions()
        rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
    finally:
        wb.close()
    if len(rows) < 4:
        return {}
    labels, days = rows[1], rows[3]
    out, current = {}, None
    for i in range(4, len(days)):
        lab = labels[i] if i < len(labels) else None
        current = lab if lab not in (None, "") else current
        d = days[i]
        if isinstance(d, (datetime, date)) and current:
            out[d.date() if isinstance(d, datetime) else d] = str(current)
    return out


def load_inputs(roles: dict[str, Path], account: str | None = None) -> Inputs:
    """account: ads_camp account_name to keep; None = the account with the most rows."""
    t = roles.get("tracker")
    sources = {}

    def table(name: str):
        """(path, sheet) for a tracker table: its own CSV if given, else the workbook sheet (None if absent)."""
        if name in roles:
            return roles[name], None
        if t is not None and TRACKER_SHEETS[name] in sheet_names(t):
            return t, TRACKER_SHEETS[name]
        return None

    def src(name: str, path: Path, sheet: str, letters: dict) -> None:
        sources[name] = {"file": Path(path).name, "sheet": sheet, "letters": letters}

    raw, letters = read_sheet_with_rows(*table("raw"))
    src("raw", table("raw")[0], table("raw")[1], letters)
    _need(raw, ["report_date", "asin", "sku", "ordered_gmv_usd", "ppc_ordered_gmv_usd", "sp_spend", "sb_spend",
                "sd_spend", "ppc_clicks", "ppc_orders"], "raw_sso_us", *table("raw"))
    raw = pd.DataFrame({
        "date": _to_date(raw["report_date"]), "asin": raw["asin"].astype(str), "sku": raw["sku"].astype(str),
        "gmv": to_num(raw["ordered_gmv_usd"]), "ads_gmv": to_num(raw["ppc_ordered_gmv_usd"]),
        "spend": to_num(raw["sp_spend"]) + to_num(raw["sb_spend"]) + to_num(raw["sd_spend"]),
        "clicks": to_num(raw["ppc_clicks"]), "orders": to_num(raw["ppc_orders"]),
        "units": to_num(raw["ordered_units"]) if "ordered_units" in raw.columns else 0.0,
        "ads_units": to_num(raw["ppc_ordered_units"]) if "ppc_ordered_units" in raw.columns else 0.0,
        "_row": raw["_row"],
    }).dropna(subset=["date"])
    if raw.empty:
        raise InputError("raw_sso_us không có dòng nào đọc được ngày ở cột report_date",
                         {"file": Path(table("raw")[0]).name, "sheet": table("raw")[1] or "raw_sso_us", "columns": ["report_date"], "row": 2})

    info, letters = read_sheet_with_rows(*table("products"))
    src("products", table("products")[0], table("products")[1], letters)
    _need(info, ["asin", "sku", "product_line_old", "status"], "Product Info", *table("products"))
    pic_cols = [c for c in info.columns if c.startswith("PIC ") or c == "PMKT PIC_Final"]
    products = pd.DataFrame({"asin": info["asin"].astype(str), "sku": info["sku"].astype(str),
                             "pl": _clean_text(info["product_line_old"]), "status": _clean_text(info["status"]),
                             "_row": info["_row"]})
    for c in ("lifecycle", "product_name"):
        if c in info.columns:
            products[c] = _clean_text(info[c])
    for c in pic_cols:
        products[c] = _clean_text(info[c])
    if "Team Label" in info.columns:  # scope for the score model
        products["team"] = _clean_text(info["Team Label"])

    k, letters = read_sheet_with_rows(*table("kpi"))
    src("kpi", table("kpi")[0], table("kpi")[1], letters)
    _need(k, ["Month", "ASIN", "GMV", "Spend_final"], "KPI", *table("kpi"))
    kpi = pd.DataFrame({"month": to_num(k["Month"]).astype(int), "asin": k["ASIN"].astype(str),
                        "gmv": to_num(k["GMV"]), "spend": to_num(k["Spend_final"]),
                        "rs": to_num(k["RS"]) if "RS" in k.columns else 0.0,
                        "stock": to_num(k["Current Stock"]) if "Current Stock" in k.columns else float("nan"),
                        "_row": k["_row"]})

    if table("ineligible"):
        ie = read_sheet(*table("ineligible"))
        ineligible = pd.DataFrame({"date": _to_date(ie.get("log_date", pd.Series(dtype=object))),
                                   "asin": ie.get("asin", pd.Series(dtype=object)).astype(str),
                                   "status": ie.get("delivery_status", pd.Series(dtype=object)),
                                   "reason": ie.get("delivery_reasons", pd.Series(dtype=object))})
    else:
        ineligible = pd.DataFrame(columns=["date", "asin", "status", "reason"])

    ads_sheet = None if is_csv(roles["ads"]) else "ads_camp"
    a, letters = read_sheet_with_rows(roles["ads"], ads_sheet)
    src("ads", roles["ads"], ads_sheet, letters)
    _need(a, ["report_date", "account_name", "campaign_type", "campaign_name", "asin",
              "spend_usd", "clicks", "sales_14d", "order_14d"], "ads_camp", roles["ads"], ads_sheet)
    if "campaign_id" not in a.columns:
        a["campaign_id"] = a["campaign_name"].map(_campaign_id)
    names = a["account_name"].astype(str).str.strip()
    accounts = names.value_counts()
    if account and account not in accounts.index:
        raise InputError(f"ads_camp không có account '{account}' – chọn lại trong Cài đặt → Tài khoản Ads")
    chosen = account or (accounts.index[0] if len(accounts) else "")
    a = a[names == chosen]
    ads = pd.DataFrame({
        "date": _to_date(a["report_date"]), "type": a["campaign_type"].astype(str),
        "campaign_id": a["campaign_id"].map(lambda v: str(int(v)) if isinstance(v, (int, float)) else str(v).strip()),
        "campaign": a["campaign_name"].astype(str), "asin": a["asin"].astype(str),
        "spend": to_num(a["spend_usd"]), "clicks": to_num(a["clicks"]),
        "sales": to_num(a["sales_14d"]), "orders": to_num(a["order_14d"]), "_row": a["_row"],
    }).dropna(subset=["date"]).reset_index(drop=True)
    ads.attrs["accounts"], ads.attrs["account"] = list(accounts.index), chosen

    log, letters = read_sheet_with_rows(roles["activity"])
    src("activity", roles["activity"], None if is_csv(roles["activity"]) else sheet_names(roles["activity"])[0], letters)
    cols = ["id", "level", "type", "action", "from_", "to", "time", "user"]
    if log.empty:
        activity = pd.DataFrame(columns=cols + ["_row"])
    else:
        activity = log.drop(columns="_row").iloc[:, :8].copy()
        activity.columns = cols[: activity.shape[1]]
        activity = activity.reindex(columns=cols)
        activity["_row"] = log["_row"].values

    defind_text = ""
    if roles.get("defind"):
        d = read_sheet(roles["defind"])
        defind_text = "\n".join(" | ".join(str(v) for v in row if v not in (None, "") and not pd.isna(v))
                                for row in [list(d.columns)] + d.values.tolist())
        defind_text = defind_text.replace("col0 | ", "")

    inventory = pd.DataFrame(columns=["date", "asin", "sku", "stock"])
    if t is not None and "Inventory AMZ" in sheet_names(t):
        inv, letters = read_sheet_with_rows(t, "Inventory AMZ")
        src("inventory", t, "Inventory AMZ", letters)
        _need(inv, ["log_date", "asin", "inventory"], "Inventory AMZ", t, "Inventory AMZ")
        inventory = pd.DataFrame({"date": _to_date(inv["log_date"]), "asin": inv["asin"].astype(str).str.strip(),
                                  "sku": inv["sku"].astype(str) if "sku" in inv.columns else "",
                                  "stock": to_num(inv["inventory"]), "_row": inv["_row"]}).dropna(subset=["date"])

    return Inputs(raw, products, kpi, ineligible, ads, activity, _event_days(t) if t is not None else {}, defind_text,
                  sources, inventory, _contributions(t) if t is not None else {})
