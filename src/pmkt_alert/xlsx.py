"""Robust xlsx reading. Exports from the team's tools carry wrong <dimension> tags,
so every read-only sheet has its dimensions reset before iterating."""
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


def sheet_names(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _read(path: Path, sheet: str | None, header_row: int):
    """(header, [(excel_row, values), ...]) with fully empty rows dropped."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        ws.reset_dimensions()
        rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    finally:
        wb.close()
    if not rows:
        return [], []
    body = [(header_row + i, r) for i, r in enumerate(rows[1:], 1) if any(v is not None and v != "" for v in r)]
    width = max(len(r) for r in [rows[0], *(r for _, r in body)])
    head = tuple(rows[0]) + (None,) * (width - len(rows[0]))
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(head)]
    return header, [(n, tuple(r[:width]) + (None,) * (width - len(r))) for n, r in body]


def read_sheet(path: Path, sheet: str | None = None, header_row: int = 1) -> pd.DataFrame:
    header, body = _read(path, sheet, header_row)
    return pd.DataFrame([r for _, r in body], columns=header) if header else pd.DataFrame()


def read_sheet_with_rows(path: Path, sheet: str | None = None, header_row: int = 1) -> tuple[pd.DataFrame, dict]:
    """Like read_sheet plus a `_row` column (Excel row number) and {header: column letter} – for evidence."""
    header, body = _read(path, sheet, header_row)
    if not header:
        return pd.DataFrame(), {}
    df = pd.DataFrame([r for _, r in body], columns=header)
    df["_row"] = [n for n, _ in body]
    return df, {h: get_column_letter(i + 1) for i, h in enumerate(header) if not h.startswith("col")}


def to_num(s: pd.Series) -> pd.Series:
    cleaned = s.map(lambda v: v.replace(",", "") if isinstance(v, str) else v)
    return pd.to_numeric(cleaned, errors="coerce").fillna(0.0).astype(float)
