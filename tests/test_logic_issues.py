import pytest

from factory import make_logic, make_logic_v3
from pmkt_alert.logic import DEFAULT_INV, DEFAULT_PARAMS, load_logic


def _codes(lg):
    return {(i["code"], i.get("key")) for i in lg.issues}


def test_v1_file_has_no_issues(tmp_path):
    lg = load_logic(make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx"))
    assert lg.issues == [] and lg.params == {**DEFAULT_PARAMS, **DEFAULT_INV}  # inventory guard: silent defaults


def test_v3_score_model_detected_and_parsed(tmp_path):
    lg = load_logic(make_logic_v3(tmp_path / "Alert_Logic_System_Team_v3.xlsx"))
    assert lg.version == 3 and lg.model == "score_v3"
    assert lg.params["trend_acos"] == 0.25 and lg.params["camp_cut"] == 1.2 and lg.k_by_type["SPS_DISPLAY"] == 5
    assert lg.params["w1"] == 3 and lg.params["w3"] == 2 and lg.params["sens_gmv"] == 0.3
    assert lg.params["score_sd"] == 50 and lg.params["score_up"] == -25 and lg.params["sd_room"] == 15
    assert lg.params["sd_over"] == -10 and lg.params["min_time_gone"] == 0.15 and lg.params["min_spend"] == 50
    assert lg.texts["team_label"] == "Team X"
    assert [i for i in lg.issues if i["code"] == "MISSING"] == []  # Ver 1 pacing keys are not required here
    unused = [i["label"] for i in lg.issues if i["code"] == "UNUSED"]
    assert unused == []  # Department block is computed by the app, not read


def test_v1_model_detected(tmp_path):
    assert load_logic(make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx")).model == "rules_v1"


def test_v3_missing_weight_reported(tmp_path):
    import openpyxl
    p = make_logic_v3(tmp_path / "v3.xlsx")
    wb = openpyxl.load_workbook(p)
    for r in wb["Config"].iter_rows(min_col=2, max_col=3):
        if r[0].value and str(r[0].value).startswith("w4"):
            r[0].value = "trọng số khác"
    wb.save(p)
    lg = load_logic(p)
    assert lg.model == "score_v3" and any(i["code"] == "MISSING" and i["key"] == "w4" for i in lg.issues)


def test_blank_value_falls_back_and_is_reported(tmp_path):
    lg = load_logic(make_logic_v3(tmp_path / "v3.xlsx", blank="Sàn click tối thiểu"))
    issue = next(i for i in lg.issues if i.get("key") == "min_clicks")
    assert issue["code"] == "BLANK" and issue["cell"].startswith("Config!C") and lg.params["min_clicks"] == 15


def test_fallback_prefers_previous_good_logic(tmp_path):
    prev = load_logic(make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx"))
    prev.params["min_clicks"] = 20
    lg = load_logic(make_logic_v3(tmp_path / "v3.xlsx", blank="Sàn click tối thiểu"), fallback=prev)
    assert lg.params["min_clicks"] == 20
    assert next(i for i in lg.issues if i.get("key") == "min_clicks")["used_from"] == "Alert_Logic_Ver 1.xlsx"


def test_bad_matrix_label_and_range_checks(tmp_path):
    p = make_logic_v3(tmp_path / "v3.xlsx", bad_matrix=True)
    lg = load_logic(p)
    bad = next(i for i in lg.issues if i["code"] == "MATRIX_LABEL")
    assert bad["key"] == "matrix:U|U|U" and bad["cell"].startswith("Matrix_PL!F") and lg.matrix["U|U|U"][0] == "REVIEW ADS"


def test_out_of_range_reported(tmp_path):
    import openpyxl
    p = make_logic_v3(tmp_path / "v3.xlsx")
    wb = openpyxl.load_workbook(p)
    for r in wb["Config"].iter_rows(min_col=2, max_col=3):
        if r[0].value and str(r[0].value).startswith("MAINTAIN/PUSH hoặc SCALE UP"):
            r[1].value = 5  # above the CUT threshold: makes no sense
    wb.save(p)
    lg = load_logic(p)
    assert ("ORDER", "camp_good") in _codes(lg)


def test_unreadable_file_falls_back_entirely(tmp_path):
    from factory import write_xlsx
    p = write_xlsx(tmp_path / "Alert_Logic_v9.xlsx", {"Sheet": [["x"]]})
    lg = load_logic(p)
    assert lg.params == {**DEFAULT_PARAMS, **DEFAULT_INV} and any(i["code"] == "NO_CONFIG" for i in lg.issues)
