import pytest

from factory import make_logic
from pmkt_alert.logic import LogicError, load_logic, version_from_name, write_logic_version
from pmkt_alert.xlsx import read_sheet


def test_loads_params_by_label_not_position(tmp_path):
    p = make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx", insert_blank_rows_at=10)
    lg = load_logic(p)
    assert lg.version == 1
    assert lg.params["camp_cut"] == 1.15 and lg.params["trend_acos"] == 0.2
    assert lg.params["camp_optimize"] == 1.3 and lg.params["spend_pace_low"] == 0.9
    assert lg.k_by_type == {"SPS_PRODUCTS": 3, "SPS_BRANDS": 3, "SPS_DISPLAY": 5}
    assert len(lg.matrix) == 27 and lg.matrix["F|F|F"][0] == "" and lg.matrix["D|D|D"][0] == "SCALE"
    assert "pacing" in lg.tiers_text


def test_missing_param_raises(tmp_path):
    p = make_logic(tmp_path / "x.xlsx", drop="min_clicks")
    with pytest.raises(LogicError, match="min_clicks"):
        load_logic(p)


def test_version_from_name():
    assert version_from_name("Alert_Logic_Ver 12.xlsx") == 12 and version_from_name("logic.xlsx") == 1


def test_write_version(tmp_path):
    src = make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx")
    out = write_logic_version(load_logic(src), {"camp_cut": 1.25, "matrix:U|F|F": "REVIEW ADS"}, "CUT 1.15→1.25", tmp_path / "out")
    assert out.name == "Alert_Logic_Ver 2.xlsx"
    new = load_logic(out)
    assert new.params["camp_cut"] == 1.25 and new.matrix["U|F|F"][0] == "REVIEW ADS" and new.version == 2
    assert load_logic(src).params["camp_cut"] == 1.15
    assert "CUT 1.15→1.25" in str(read_sheet(out, "README_Flow", header_row=1).to_dict())


def test_valid_change(tmp_path):
    from pmkt_alert.logic import valid_change
    lg = load_logic(make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx"))
    assert valid_change("camp_cut", 1.25, lg) and not valid_change("camp_cut", "x", lg)
    assert valid_change("matrix:F|F|F", "REVIEW ADS", lg) and not valid_change("matrix:Z", "X", lg)
    assert not valid_change("nope", 1, lg)


def test_new_version_never_overwrites(tmp_path):
    src = make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx")
    out_dir = tmp_path / "out"
    v2 = write_logic_version(load_logic(src), {"camp_cut": 1.5}, "a", out_dir)
    v3 = write_logic_version(load_logic(src), {"camp_cut": 2.0}, "b", out_dir)  # from Ver 1 again
    assert v2.name == "Alert_Logic_Ver 2.xlsx" and v3.name == "Alert_Logic_Ver 3.xlsx"
    assert load_logic(v2).params["camp_cut"] == 1.5


def test_missing_readme_sheet_created(tmp_path):
    import openpyxl
    src = make_logic(tmp_path / "Alert_Logic_Ver 1.xlsx")
    wb = openpyxl.load_workbook(src); del wb["README_Flow"]; wb.save(src)
    out = write_logic_version(load_logic(src), {"camp_cut": 1.2}, "n", tmp_path / "o")
    assert "README_Flow" in openpyxl.load_workbook(out).sheetnames
